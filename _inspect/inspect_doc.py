#!/usr/bin/env python3
"""Deep inspection of the FH Burgenland Ansuchen docx: paragraphs, runs,
styles, sections, tables, images — dumped as JSON for planning purposes."""
import json, sys
from docx import Document
from docx.shared import Emu as EMU

DOC = "/home/cloude/from-word-to-latex/01_Ansuchen Zuteilung wiss Arbeit.docx"

doc = Document(DOC)

def emu_to_cm(v):
    return round(EMU(v).cm, 2) if v is not None else None

out = {"sections": [], "paragraphs": [], "tables": []}

# --- sections / page setup ---
for i, s in enumerate(doc.sections):
    out["sections"].append({
        "index": i,
        "page_width_cm": emu_to_cm(s.page_width),
        "page_height_cm": emu_to_cm(s.page_height),
        "left_margin_cm": emu_to_cm(s.left_margin),
        "right_margin_cm": emu_to_cm(s.right_margin),
        "top_margin_cm": emu_to_cm(s.top_margin),
        "bottom_margin_cm": emu_to_cm(s.bottom_margin),
        "header_distance_cm": emu_to_cm(s.header_distance),
        "footer_distance_cm": emu_to_cm(s.footer_distance),
        "orientation": str(s.orientation),
        "different_first_page_header": s.different_first_page_header_footer,
        "header_text": "\n".join(p.text for p in s.header.paragraphs if p.text.strip()),
        "footer_text": "\n".join(p.text for p in s.footer.paragraphs if p.text.strip()),
        "first_page_header_text": "\n".join(p.text for p in s.first_page_header.paragraphs if p.text.strip()),
        "first_page_footer_text": "\n".join(p.text for p in s.first_page_footer.paragraphs if p.text.strip()),
    })

# --- paragraphs (body) ---
for i, p in enumerate(doc.paragraphs):
    runs = []
    for r in p.runs:
        rd = {}
        if r.text:
            rd["text"] = r.text
        if r.bold: rd["bold"] = True
        if r.italic: rd["italic"] = True
        if r.underline: rd["underline"] = True
        if r.font.size: rd["size_pt"] = r.font.size.pt
        if r.font.name: rd["font"] = r.font.name
        if r.font.cs_bold: rd["cs_bold"] = True
        if r.font.all_caps: rd["all_caps"] = True
        if r.font.small_caps: rd["small_caps"] = True
        if rd.get("text"):
            runs.append(rd)
    d = {"i": i, "style": p.style.name}
    # paragraph-level direct formatting
    pf = p.paragraph_format
    if pf.alignment is not None: d["align"] = str(pf.alignment)
    if pf.left_indent is not None: d["l_indent_cm"] = emu_to_cm(pf.left_indent)
    if pf.right_indent is not None: d["r_indent_cm"] = emu_to_cm(pf.right_indent)
    if pf.first_line_indent is not None: d["fl_indent_cm"] = emu_to_cm(pf.first_line_indent)
    if pf.space_before is not None: d["space_before_pt"] = pf.space_before.pt
    if pf.space_after is not None: d["space_after_pt"] = pf.space_after.pt
    if pf.line_spacing is not None: d["line_spacing"] = str(pf.line_spacing)
    if p.text.strip():
        d["text"] = p.text
    if runs:
        d["runs"] = runs
    # inline images
    if p._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline'):
        d["has_image"] = True
    out["paragraphs"].append(d)

# --- tables ---
for ti, t in enumerate(doc.tables):
    rows = []
    for row in t.rows:
        rows.append([c.text for c in row.cells])
    out["tables"].append({
        "index": ti,
        "style": t.style.name if t.style else None,
        "n_rows": len(t.rows),
        "n_cols": len(t.columns),
        "grid": rows,
    })

json.dump(out, sys.stdout, ensure_ascii=False, indent=1)